#!/usr/bin/env python3
"""Build every image, logo and video the Canton Fair kiosk ships with.

Sources are the client's raw folder `canton fair/` and the approved
`HOSUN_spreads_v4.pdf`, both at the repository root and both untracked
(1.1 GB). Outputs land in `fair/assets/` and are committed, so the site
itself never needs this script — run it only when the source material changes.

    python3 -m venv .venv && .venv/bin/pip install Pillow pillow-heif python-pptx opencv-python-headless
    .venv/bin/python fair/tools/build_assets.py            # skips the film if it exists
    .venv/bin/python fair/tools/build_assets.py --film     # re-encode the film too
    .venv/bin/python fair/tools/build_precache.py          # always run afterwards

Needs `ffmpeg`, `pdfimages` and `pdftocairo` (Homebrew: ffmpeg, poppler).
"""
import argparse
import io
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from PIL import Image, ImageOps
import pillow_heif

pillow_heif.register_heif_opener()  # two of the factory photos are HEIC files named .jpg
Image.MAX_IMAGE_PIXELS = None

REPO = Path(__file__).resolve().parents[2]
SRC = REPO / "canton fair"
SPREADS = REPO / "HOSUN_spreads_v4.pdf"
OUT = REPO / "fair" / "assets"
FILM = SRC / "恒而思达中文版2026.m4v"          # stills are taken from this one
FILMS = {"film-en-1080.mp4": SRC / "恒而思达英文版2026.m4v", "film-zh-1080.mp4": FILM}

# ---- photography: full-bleed slides, 1920 px on the long edge --------------
PHOTOS = {
    "factory/stamping-robot": "工厂展厅现场照片/冲压机器人.png",
    "factory/welding-robot": "工厂展厅现场照片/焊接机器人.jpg",
    "factory/welding": "工厂展厅现场照片/焊接.jpg",
    "factory/machining": "工厂展厅现场照片/机加工.jpg",
    "factory/robot-arm-1": "工厂展厅现场照片/机械臂1.jpg",
    "factory/robot-arm-2": "工厂展厅现场照片/机械臂2.jpg",
    "factory/painting-line": "工厂展厅现场照片/喷涂线.JPG",
    "factory/painting": "工厂展厅现场照片/喷涂1.jpg",
    "factory/assembly-1": "工厂展厅现场照片/组装1.png",
    "factory/assembly-2": "工厂展厅现场照片/组转2.JPG",
    "factory/assembly-3": "工厂展厅现场照片/组装3.JPG",
    "factory/lifting-columns": "工厂展厅现场照片/立柱.jpg",
    "factory/c-arm-line": "工厂展厅现场照片/C-arm.JPG",
    "factory/imaging-system": "工厂展厅现场照片/影像设备.png",
    "factory/showroom-1": "工厂展厅现场照片/展厅1.JPG",
    "factory/showroom-2": "工厂展厅现场照片/展厅2.JPG",
    "factory/history-wall": "工厂展厅现场照片/发展历程.JPG",
    "factory/partner-wall": "工厂展厅现场照片/合作伙伴.JPG",
    "factory/team": "工厂展厅现场照片/门口合影.JPG",
    "factory/new-campus": "工厂展厅现场照片/新厂效果图.jpg",
    "factory/new-campus-2": "工厂展厅现场照片/新厂规划设计2.jpg",
    "factory/new-campus-3": "工厂展厅现场照片/新厂规划设计5.jpg",
}

# ---- stills from the 4K brand film -----------------------------------------
# (seconds, keep) — `keep` is the share of frame height kept from the top, so
# the burned-in bilingual subtitles (and the lower-left chapter labels) are cut.
# `keep` may instead be a (left, top, right, bottom) box in frame fractions.
FRAMES = {
    "film/sunrise": (4, 0.80),
    "film/home-care": (16, 0.82),
    # the whole team at the table: the crop stops above the subtitle and the
    # "研发设计 R&D Design" chapter label is painted out (see unlabel)
    "film/rnd": (46.6, (0.0, 0.28, 1.0, 0.835)),
    "film/precision": (52, 0.70),
    "film/robot": (56, 0.70),
    "film/assembly": (108, 0.82),
    "film/bed-studio-1": (120, 0.82),
    "film/bed-studio-2": (124, 0.82),
    "film/bed-studio-3": (128, 0.82),
    "film/c-arm": (136, 0.82),
    "film/xray-detector": (140, 0.82),
    "film/accessible-kitchen": (144, 0.82),
    "film/ankle-trainer": (148, 0.82),
}

# ---- product shots: mostly on white, 1200 px --------------------------------
P = "产品/"
D = P + "电动升降桌/升降桌系列/"
PRODUCTS = {
    # nursing beds
    "product/em-2001-a": P + "床/EM-2001 (1).jpg",
    "product/em-2001-b": P + "床/EM-2001 (2).jpg",
    "product/em-2002-a": P + "床/EM-2002.jpg",
    "product/em-2002-b": P + "床/EM-2002 (2).jpg",
    "product/em-3001-a": P + "床/EM-3001 (1).jpg",
    "product/em-3001-b": P + "床/EM-3001 (2).jpg",
    "product/em-3002-a": P + "床/EM-3002 (1).jpg",
    "product/em-3002-b": P + "床/EM-3002 (2).jpg",
    "product/em-3003-a": P + "床/EM-3003.jpg",
    "product/em-3003-b": P + "床/EM-3003 (2).jpg",
    "product/em-5001-a": P + "床/EM-5001.jpg",
    "product/em-5001-b": P + "床/EM-5001(2).jpg",
    "product/em-5002-a": P + "床/EM-5002.jpg",
    "product/em-5002-b": P + "床/EM-5002(2）.jpg",
    # powered rollator RW-01 (the folder names are the client's; the bed-named file is the rollator)
    "product/rollator": P + "助行器/74898422ee450650e3534d1771d44ea3.jpg",
    "product/rollator-b": P + "助行器/1.28 康佳乐款电动护理床1535.jpg",
    "product/rollator-folded": P + "助行器/折叠.jpg",
    "product/rollator-dims": P + "助行器/2 (2).jpg",
    "product/rollator-seat": P + "助行器/2 (1).jpg",
    "product/rollator-stop": P + "助行器/1234 (1).jpg",
    "product/rollator-straight": P + "助行器/1234 (2).jpg",
    "product/rollator-downhill": P + "助行器/1234 (3).jpg",
    "product/rollator-uphill": P + "助行器/1234 (4).jpg",
    # transfer & lifting
    "product/transfer-robot": "../1.png",
    "product/xray-c-arm": "../C-arm-xray.jpg",
    "product/transfer-device-single": P + "中车移位机器/single.jpg",
    "product/transfer-device": P + "中车移位机器/121.jpg",
    "product/transfer-device-b": P + "中车移位机器/图片_副本.jpg",
    "product/transfer-device-views": P + "中车移位机器/康复助站设备三视图展示.png",
    "product/hoist": P + "升降式移位机/吊装移位机.jpg",
    "product/hoist-room": P + "升降式移位机/图片生成与修改 (1).png",
    "product/hoist-b": P + "升降式移位机/图片生成与修改.png",
    # bath & toilet
    "product/bath-lift-a": P + "升降洗澡椅/1 (2).jpg",
    "product/bath-lift-b": P + "升降洗澡椅/1 (1).jpg",
    "product/bath-lift-c": P + "升降洗澡椅/1 (3).jpg",
    "product/bath-lift-views": P + "升降洗澡椅/123.jpg",
    "product/toilet-lift": P + "升降马桶椅/Gemini_Generated_Image_tkc1ootkc1ootkc1.jpg",
    "product/toilet-lift-b": P + "升降马桶椅/Gemini_Generated_Image_uwzhrxuwzhrxuwzh.jpg",
    "product/shower-chair": P + "洗澡椅/洗澡1椅 - 副本_副本.png",
    "product/shower-chair-scene": P + "洗澡椅/1234.jpg",
    "product/shower-chair-dims": P + "洗澡椅/Gemini_Generated_Image_2wka212wka212wka_副本.jpg",
    "product/basin": P + "电动升降洗手台盆/sink.jpg",
    "product/basin-scene": P + "电动升降洗手台盆/1 (1).jpg",
    "product/basin-views": P + "电动升降洗手台盆/1 (2).jpg",
    "product/basin-models": P + "电动升降洗手台盆/1 (3).jpg",
    # accessible kitchen & home
    "product/sink": P + "电动升降水槽/1.jpg",
    "product/sink-kitchen": P + "电动升降水槽/轮椅车轮清洗机RM2作业流程与效益对比.png",
    "product/hob": P + "电动升降灶具/升降灶具.jpg",
    "product/hob-b": P + "电动升降灶具/轮椅车轮清洗机RM2作业流程与效益对比 (1).png",
    "product/cabinet-lift": P + "升降橱柜/电动升降厨柜.jpg",
    "product/cabinet-lift-spices": P + "升降橱柜/5034e75eea2cd37ae0b0c10610235982.jpg",
    "product/cabinet-lift-glass": P + "升降橱柜/WPS图片.png",
    "product/cabinet-lift-photo": P + "升降橱柜/微信图片_20250606134738.jpg",
    "product/wardrobe-rail": P + "升降衣杆/sg-11134207-7rdws-md19pdd029u634_副本.jpg",
    "product/wardrobe-scene-1": P + "升降衣杆/WPS图片(1).png",
    "product/wardrobe-scene-2": P + "升降衣杆/WPS图片.png",
    # rehab & facility
    "product/ankle-a": P + "踝关节康复训练器/1 (1).jpg",
    "product/ankle-b": P + "踝关节康复训练器/1 (2).jpg",
    "product/ankle-c": P + "踝关节康复训练器/1 (3).jpg",
    "product/ankle-d": P + "踝关节康复训练器/1 (4).jpg",
    "product/ankle-scene": P + "踝关节康复训练器/ScreenShot_2026-10-16_162502_942.png",
    "product/wheel-washer": P + "轮椅清洗机/567.jpg",
    "product/wheel-washer-b": P + "轮椅清洗机/23.jpg",
    "product/wheel-washer-c": P + "轮椅清洗机/345.jpg",
    # desks & tables
    "product/desk-office-dual": D + "2..png",
    "product/desk-office-dual-b": D + "2-2.png",
    "product/desk-office-wood": D + "3.jpg",
    "product/desk-office-wood-views": D + "3-2.jpg",
    "product/desk-coffee": D + "1.jpg",
    "product/desk-coffee-views": D + "1-1.jpg",
    "product/desk-four-leg": D + "5.png",
    "product/desk-four-leg-views": D + "5.jpg",
    "product/desk-executive": D + "升降经理桌.png",
    "product/desk-gaming": D + "4 (1).jpg",
    "product/desk-gaming-views": D + "4 (2).jpg",
    "product/desk-student": D + "7 (2).png",
    "product/desk-student-views": D + "7 (1).png",
    "product/desk-health": D + "智能检测桌.png",
    "product/table-tilting": D + "6 (1).jpg",
    "product/table-tilting-views": D + "6 (2).jpg",
    "product/table-pneumatic": D + "气动1 (1).jpg",
    "product/table-angled": D + "气动1 (2).jpg",
    "product/table-side": D + "气动1 (3).jpg",
    "product/table-side-round": D + "气动1 (4).jpg",
    "product/table-overbed": D + "气动1 (5).jpg",
}

# ---- certificate scans ---------------------------------------------------------
# Folder scans; the English set from slide 7 of 恒升介绍英文版.pptx is in DECK_CERTS.
CERTS = {
    "cert/iso-13485": "认证/ISO 13485_01.png",   # slide 7's copy expired 2020-12-18; this one runs to 2026-12-18
    "cert/fcc-sdoc": "认证/FCC_01.png",
}
DECK = SRC / "物料/恒升介绍英文版.pptx"
# shape id on slide 7 → output name
DECK_CERTS = {7: "cert/iatf-16949", 8: "cert/iso-9001", 5: "cert/iso-14001", 6: "cert/iso-45001", 9: "cert/sa8000"}
# slide 8: the ISO 3834 certificate, then the eight welders' certificates left to right
DECK_WELDING = {21: "cert/iso-3834", 18: "cert/welder-1", 14: "cert/welder-2", 12: "cert/welder-3", 17: "cert/welder-4",
                3: "cert/welder-5", 13: "cert/welder-6", 11: "cert/welder-7", 10: "cert/welder-8"}

# ---- partner logos, full colour -------------------------------------------------
# Order follows slide 6 of the company deck, row by row; logos from the newer
# `canton fair/logo/` folder join the row of their kind. Sources: a file in that
# folder, a picture on slide 6 ("deck:<shape id>"), or the spread's Toyota mask
# recoloured in Toyota red (the deck only has a 151 px copy).
L = "logo/"
PARTNERS = [
    # industrial — Japan, Germany, US
    ("daikin", L + "Daikin-Logo.png"), ("toyota", "spread-toyota"), ("viessmann", L + "viessmann-logo.svg"),
    ("bmw", L + "BMW.svg"), ("zf", L + "ZF-Group_Logo_Blue_RGB.png"), ("mitsubishi-electric", L + "Mitsubishi_Electric_logo.svg.png"),
    ("emerson", L + "Emerson-Electric-Logo.png"), ("assa-abloy", L + "Assa_Abloy.svg.png"), ("wolf", L + "WOLF_Logo.png"),
    ("aichi", L + "dfytjyb-f563198e-8fe3-4b10-8c37-f3b32aac5b9c.png"), ("fujikoki", L + "fujikoki_czech_s_r_o_cover.jpeg"),
    # medical imaging and devices
    ("philips", L + "Phillips-Logo-2008.png"), ("siemens", L + "Siemens_AG_logo.svg.png"), ("maquet", L + "MAQUET_AG_Logo.png"),
    ("carestream", "deck:37"), ("nuctech", L + "0bfbf27e697240cb863b0234fe15758d.webp"), ("meyer", L + "170468468514657300.svg"),
    ("perlove", L + "2022090109501942569831071.webp"), ("getein", L + "2021120911073720537.jpg"),
    # care equipment
    ("arjo", L + "Arjo_Logotype_Tagline_RGB.jpg"), ("paramount-bed", L + "paramound.png"), ("barco", L + "barco.png"),
    ("vishee", L + "10069990.jpg"),
    # lifting systems
    ("kessebohmer", L + "kesseboehmer_logo.png"), ("dewertokin", L + "logo-dewert-okin.svg"), ("linak", "deck:28"),
    # health and senior-care groups
    ("suzhou-health", L + "苏州康养集团logo.png"), ("genertec", L + "logo.png"), ("taikang", L + "1650351975-截屏2022-04-19-下午3.05.58.png"),
    # 南通鑫汇养老 left out on request, like Hubei
    # software
    ("jbinfo", L + "ae99f18f2c70beb4ac80cc760ff5ad40.png"),
]

# ---- brand marks, vector, from the two VI manuals --------------------------
# (pdf, page, viewBox in PDF points measured off a 72 dpi render)
BRAND = {
    "hosun": ("物料/恒升vi.pdf", 7, "455 301 390 180"),
    "hosun-mark": ("物料/恒升vi.pdf", 7, "455 301 160 180"),
    "herstar": ("物料/恒而思达vi20230424.pdf", 5, "1017 312 533 204"),
}


def save_webp(im, dest, quality):
    dest.parent.mkdir(parents=True, exist_ok=True)
    im.save(dest, "WEBP", quality=quality, method=6)
    return dest


def load(path, background=(255, 255, 255)):
    im = ImageOps.exif_transpose(Image.open(path))
    if im.mode in ("RGBA", "LA", "P"):
        im = im.convert("RGBA")
        flat = Image.new("RGB", im.size, background)
        flat.paste(im, mask=im.getchannel("A"))
        return flat
    return im.convert("RGB")


def fit(im, long_edge):
    im = im.copy()
    im.thumbnail((long_edge, long_edge), Image.LANCZOS)
    return im


def build_rasters():
    for name, rel in PHOTOS.items():
        save_webp(fit(load(SRC / rel), 1920), OUT / "img" / f"{name}.webp", 76)
    for name, rel in PRODUCTS.items():
        im = load(SRC / rel)
        # shots on white lose their uneven borders, so every catalogue thumbnail sits the same way
        corners = im.convert("L").resize((8, 8))
        if min(corners.getpixel(xy) for xy in ((0, 0), (7, 0), (0, 7), (7, 7))) > 238:
            im = trim_white(im)
        save_webp(fit(im, 1200), OUT / "img" / f"{name}.webp", 80)
    build_certs()
    build_cover()
    build_qr()
    print(f"rasters: {len(PHOTOS)} photos, {len(PRODUCTS)} products")


def unlabel(png):
    """Inpaint the film's white lower-left chapter label (needs opencv-python-headless)."""
    import cv2
    import numpy as np
    im = cv2.imread(str(png))
    h, w = im.shape[:2]
    top = int(h * 0.70)
    band = cv2.cvtColor(im[top:int(h * 0.82), :int(w * 0.35)], cv2.COLOR_BGR2GRAY)
    ys, xs = np.nonzero(band > 225)
    x0, x1, y0, y1 = xs.min() - 12, xs.max() + 12, top + ys.min() - 12, top + ys.max() + 12
    mask = np.zeros((h, w), np.uint8)
    mask[y0:y1, x0:x1] = (cv2.cvtColor(im[y0:y1, x0:x1], cv2.COLOR_BGR2GRAY) > 175).astype(np.uint8) * 255
    mask = cv2.dilate(mask, np.ones((7, 7), np.uint8))
    cv2.imwrite(str(png), cv2.inpaint(im, mask, 6, cv2.INPAINT_NS))


def build_frames(tmp):
    for name, (t, keep) in FRAMES.items():
        raw = tmp / f"{name.replace('/', '-')}.png"
        subprocess.run(["ffmpeg", "-loglevel", "error", "-y", "-ss", str(t), "-i", FILM,
                        "-frames:v", "1", raw], check=True)
        if name == "film/rnd":
            unlabel(raw)
        im = Image.open(raw).convert("RGB")
        box = keep if isinstance(keep, tuple) else (0, 0, 1, keep)
        im = im.crop(tuple(round(v * (im.width if i % 2 == 0 else im.height)) for i, v in enumerate(box)))
        save_webp(fit(im, 1920), OUT / "img" / f"{name}.webp", 78)
    # poster for both film players: the HERSTAR logo on a blurred background at 00:31
    # of the English cut — 31.15 s falls between the end of the subtitle (31.0 s) and a
    # light flare (31.5 s), so the frame carries no subtitle
    raw = tmp / "poster.png"
    subprocess.run(["ffmpeg", "-loglevel", "error", "-y", "-ss", "31.15", "-i", FILMS["film-en-1080.mp4"],
                    "-frames:v", "1", raw], check=True)
    save_webp(fit(Image.open(raw).convert("RGB"), 1920), OUT / "img" / "film" / "poster.webp", 82)
    print(f"film stills: {len(FRAMES)} + poster")


def deck_pictures(slide_no):
    """{shape id: image bytes} for every picture on one slide of the company deck."""
    from pptx import Presentation  # only needed here
    def walk(shapes):
        for sh in shapes:
            yield sh
            if sh.shape_type == 6:
                yield from walk(sh.shapes)
    slide = Presentation(DECK).slides[slide_no - 1]
    return {sh.shape_id: sh.image.blob for sh in walk(slide.shapes) if sh.shape_type == 13}


def trim_white(im, tolerance=14, margin=0.04):
    """Crop near-white borders, then add an even margin back."""
    rgb = im.convert("RGB")
    diff = ImageOps.invert(rgb.convert("L")).point(lambda v: 255 if v > tolerance else 0)
    box = diff.getbbox()
    if not box:
        return im
    pad = round(max(box[2] - box[0], box[3] - box[1]) * margin)
    box = (max(box[0] - pad, 0), max(box[1] - pad, 0), min(box[2] + pad, im.width), min(box[3] + pad, im.height))
    return im.crop(box)


def build_certs():
    for name, rel in CERTS.items():
        im = load(SRC / rel)
        im.thumbnail((1600, 2260), Image.LANCZOS)  # sharp when the viewer fills a tablet screen
        save_webp(im, OUT / "img" / f"{name}.webp", 74)
    for shape_id, name in DECK_CERTS.items():
        im = Image.open(io.BytesIO(deck_pictures(7)[shape_id])).convert("RGB")
        im.thumbnail((1200, 1700), Image.LANCZOS)
        save_webp(im, OUT / "img" / f"{name}.webp", 78)
    # slide 8 — ISO 3834-2 and the individual welder certificates, in the slide's order
    pics = deck_pictures(8)
    for shape_id, name in DECK_WELDING.items():
        im = Image.open(io.BytesIO(pics[shape_id])).convert("RGBA")
        flat = Image.new("RGB", im.size, (255, 255, 255))
        flat.paste(im, mask=im.getchannel("A"))
        save_webp(flat, OUT / "img" / f"{name}.webp", 86)
    print(f"certificates: {len(CERTS)} scans + {len(DECK_CERTS)} from slide 7 + {len(DECK_WELDING)} from slide 8")


def build_cover():
    """Slide 1 of the company deck is the site's cover; its three logo pictures, kept with alpha."""
    pics = deck_pictures(1)
    for shape_id, name, width in ((8, "hosun", 760), (9, "herstar", 760), (7, "company-name", 2000)):
        im = Image.open(io.BytesIO(pics[shape_id])).convert("RGBA")
        im.thumbnail((width, width), Image.LANCZOS)
        save_webp(im, OUT / "img" / "cover" / f"{name}.webp", 88)
    # the cover's background artwork (canton fair/产品/cover.pdf), rendered at iPad Pro width
    with tempfile.TemporaryDirectory() as tmp:
        raw = Path(tmp) / "bg"
        subprocess.run(["pdftocairo", "-png", "-singlefile", "-scale-to-x", "2800", "-scale-to-y", "-1",
                        SRC / "产品/cover.pdf", raw], check=True)
        save_webp(Image.open(f"{raw}.png").convert("RGB"), OUT / "img" / "cover" / "background.webp", 88)
    print("cover: slide 1 logos + background")


# slide 18 — the four social-media QR codes, in the slide's order
DECK_QR = {6: "qr/wechat-official", 8: "qr/wechat-channels", 11: "qr/douyin", 14: "qr/jd"}


def build_qr():
    """QR codes on white (the deck's are dark on transparent), lossless so they stay scannable."""
    pics = deck_pictures(18)
    for shape_id, name in DECK_QR.items():
        im = Image.open(io.BytesIO(pics[shape_id])).convert("RGBA")
        flat = Image.new("RGB", im.size, (255, 255, 255))
        flat.paste(im, mask=im.getchannel("A"))
        dest = OUT / "img" / f"{name}.webp"
        dest.parent.mkdir(parents=True, exist_ok=True)
        flat.save(dest, "WEBP", lossless=True)
    print(f"QR codes: {len(DECK_QR)}")


def build_partners(tmp):
    """Full-colour logos on white, trimmed to the mark, for white tiles."""
    dest_dir = OUT / "logos" / "partners"
    if dest_dir.exists():
        shutil.rmtree(dest_dir)
    dest_dir.mkdir(parents=True)
    deck = deck_pictures(6)
    subprocess.run(["pdfimages", "-f", "1", "-l", "1", "-png", SPREADS, tmp / "s"], check=True)
    for name, source in PARTNERS:
        if source.endswith(".svg"):
            shutil.copy(SRC / source, dest_dir / f"{name}.svg")  # vector, tight viewBox already
            continue
        if source == "spread-toyota":
            shape = Image.open(tmp / "s-044.png").convert("L")
            im = Image.new("RGB", shape.size, (255, 255, 255))
            im.paste(Image.new("RGB", shape.size, (235, 10, 30)), mask=shape)
        elif source.startswith("deck:"):
            im = load(io.BytesIO(deck[int(source[5:])]))
        else:
            im = load(SRC / source)
        im = trim_white(im, tolerance=24)
        im.thumbnail((720, 240), Image.LANCZOS)
        im.save(dest_dir / f"{name}.png", optimize=True)
    print(f"partner logos: {len(PARTNERS)}")


INK = re.compile(r'fill="rgb\((1[0-9]|2[0-9])\.\d+%, (\d|1\d|2\d)\.\d+%, (\d|1\d|2\d)\.\d+%\)"')


def build_brand(tmp):
    dest_dir = OUT / "logos"
    for name, (rel, page, box) in BRAND.items():
        svg = tmp / f"{name}.svg"
        subprocess.run(["pdftocairo", "-svg", "-f", str(page), "-l", str(page), SRC / rel, svg],
                       check=True, stderr=subprocess.DEVNULL)
        text = svg.read_text()
        w, h = box.split()[2:]
        text = re.sub(r'<svg ([^>]*?)width="[^"]*" height="[^"]*" viewBox="[^"]*"',
                      lambda m: f'<svg {m.group(1)}width="{w}" height="{h}" viewBox="{box}"',
                      text, count=1)
        # the manual's page background is a white rectangle — drop it so the mark sits on any ground
        text = text.replace('fill="rgb(100%, 100%, 100%)"', 'fill="none"')
        (dest_dir / f"{name}.svg").write_text(text)
        # a bone-white variant for the night skin: only the near-black ink changes, the sun keeps its gradient
        (dest_dir / f"{name}-light.svg").write_text(re.sub(INK, 'fill="#f3ece2"', text))
    print(f"brand marks: {len(BRAND)} × dark + light")


def build_icons(tmp):
    """Home-screen icons: the H mark on the cover's peach, rendered from the VI manual."""
    rel, page, box = BRAND["hosun-mark"]
    x, y, w, h = (float(v) for v in box.split())
    raw = tmp / "mark"
    subprocess.run(["pdftocairo", "-png", "-singlefile", "-r", "400", "-f", str(page), "-l", str(page),
                    "-x", str(round(x * 400 / 72)), "-y", str(round(y * 400 / 72)),
                    "-W", str(round(w * 400 / 72)), "-H", str(round(h * 400 / 72)), SRC / rel, raw], check=True)
    mark = Image.open(f"{raw}.png").convert("RGB")
    dest_dir = OUT / "icons"
    dest_dir.mkdir(parents=True, exist_ok=True)
    for size in (512, 192, 180):
        ground = Image.new("RGB", (size, size), (247, 226, 193))
        m = mark.copy()
        m.thumbnail((round(size * 0.62), round(size * 0.62)), Image.LANCZOS)
        # the manual's page is white; key it out against the peach ground
        alpha = ImageOps.invert(m.convert("L")).point(lambda v: min(255, v * 3))
        ground.paste(m, ((size - m.width) // 2, (size - m.height) // 2), alpha)
        ground.save(dest_dir / f"icon-{size}.png", optimize=True)
    print("icons: 3")


def build_film(only=None):
    # H.264 high profile, so every tablet browser plays it; capped so each
    # version stays near 45 MB inside the offline cache
    for name, source in FILMS.items():
        if only and name != only:
            continue
        dest = OUT / "video" / name
        dest.parent.mkdir(parents=True, exist_ok=True)
        subprocess.run([
            "ffmpeg", "-loglevel", "error", "-y", "-i", source,
            "-vf", "scale=1920:-2:flags=lanczos", "-c:v", "libx264", "-preset", "slow",
            "-crf", "24", "-maxrate", "3000k", "-bufsize", "6000k", "-profile:v", "high",
            "-pix_fmt", "yuv420p", "-c:a", "aac", "-b:a", "128k", "-movflags", "+faststart", dest,
        ], check=True)
        print(f"film {name}: {dest.stat().st_size / 1e6:.1f} MB")


def main():
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("--film", action="store_true", help="re-encode the brand film")
    parser.add_argument("--logos-only", action="store_true", help="rebuild partner logos and brand marks only")
    args = parser.parse_args()
    missing = [p for p in (SRC, SPREADS, *FILMS.values()) if not p.exists()]
    if missing:
        sys.exit(f"missing source material: {', '.join(map(str, missing))}")
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        if args.logos_only:
            build_partners(tmp)
            build_brand(tmp)
            build_icons(tmp)
            return
        build_rasters()
        build_frames(tmp)
        build_partners(tmp)
        build_brand(tmp)
        build_icons(tmp)
    for name in FILMS:
        if args.film or not (OUT / "video" / name).exists():
            build_film(only=name)


if __name__ == "__main__":
    main()
